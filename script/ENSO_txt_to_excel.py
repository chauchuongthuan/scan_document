import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# OUTGOING LONG WAVE RADIATION EQUATOR & SEA LEVEL PRESS (sửa tên file)
def convert_txt_to_sea_level_press_excel(input_txt_path, output_excel_path='C:/Users/COMPUTER/Desktop/OUTGOING LONG WAVE RADIATION EQUATOR.xlsx'):
    # 1. Đọc nội dung file txt
    with open(input_txt_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    columns = ['YEAR', 'JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN', 'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC']
    
    anomaly_rows = []
    original_rows = []
    standardized_rows = []
    current_section = None
    
    # 2. Phân tích nội dung theo từng dòng
    for line in lines:
        line_upper = line.upper()
        
        # Nhận diện phần dữ liệu
        if 'ORIGINAL' in line_upper:
            current_section = 'ORIGINAL DATA'
            continue
        elif 'ANOMALY' in line_upper:
            current_section = 'ANOMALY'
            continue
        elif 'STANDARDIZED' in line_upper:
            current_section = 'STANDARDIZED'
            continue
            
        # Trích xuất tất cả các số trong dòng (bao gồm số âm liền kề như -4.0-999.9)
        tokens = re.findall(r'-?\d+\.?\d*', line)
        
        # Một dòng hợp lệ gồm 1 YEAR + 12 tháng = 13 giá trị
        if tokens and len(tokens) == 13 and len(tokens[0]) == 4 and tokens[0].isdigit():
            year = int(tokens[0])
            if 1900 <= year <= 2100:
                row_vals = [year] + [float(x) for x in tokens[1:]]
                
                if current_section == 'ANOMALY':
                    anomaly_rows.append(row_vals)
                elif current_section == 'ORIGINAL DATA':
                    original_rows.append(row_vals)
                elif current_section == 'STANDARDIZED':
                    standardized_rows.append(row_vals)
                    
    # Tạo DataFrame
    df_anomaly = pd.DataFrame(anomaly_rows, columns=columns)
    df_original = pd.DataFrame(original_rows, columns=columns)
    df_standardized = pd.DataFrame(standardized_rows, columns=columns)
    
    # 3. Ghi vào file Excel và định dạng đẹp mắt bằng openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Xóa sheet mặc định
    
    # Định dạng style
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Xanh navy đậm
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    sheets_data = [
        ('ORIGINAL DATA', df_original),
        ('ANOMALY', df_anomaly),
        ('STANDARDIZED', df_standardized)
    ]
    
    for sheet_name, df in sheets_data:
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True  # Hiển thị đường lưới
        
        # Ghi Header
        ws.append(columns)
        for col_num in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Ghi các dòng dữ liệu
        for r_idx, row in df.iterrows():
            ws.append(list(row))
            row_num = r_idx + 2
            
            # Cột YEAR
            year_cell = ws.cell(row=row_num, column=1)
            year_cell.font = Font(name="Segoe UI", size=10, bold=True)
            year_cell.alignment = Alignment(horizontal="center", vertical="center")
            year_cell.number_format = '0'
            year_cell.border = thin_border
            
            # Các cột JAN -> DEC
            for col_num in range(2, 14):
                cell = ws.cell(row=row_num, column=col_num)
                cell.font = data_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '0.0'
                cell.border = thin_border
                
        # Cố định hàng đầu tiên (Freeze Pane)
        ws.freeze_panes = 'A2'
        
        # Chỉnh độ rộng cột
        ws.column_dimensions['A'].width = 10
        for col_idx in range(2, 14):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 9
            
    # Lưu file
    wb.save(output_excel_path)
    print(f"Đã tạo thành công file Excel: {output_excel_path}")

def process_relative_oni_to_excel(input_data, output_filename="C:/Users/COMPUTER/Desktop/Relative Oceanic Nino Index.xlsx"):
    """
    Đọc dữ liệu Relative Oceanic Nino Index từ đường dẫn file txt hoặc chuỗi text,
    xuất ra file Excel chuẩn định dạng.
    """
    # 1. Đọc dữ liệu từ file path (.txt) hoặc từ chuỗi string truyền vào
    if isinstance(input_data, str):
        if input_data.strip().endswith(".txt"):
            df = pd.read_csv(input_data, sep=r'\s+', engine='python')
        else:
            df = pd.read_csv(io.StringIO(input_data.strip()), sep=r'\s+', engine='python')
    else:
        df = pd.read_csv(input_data, sep=r'\s+', engine='python')

    # Đảm bảo tên cột và kiểu dữ liệu chuẩn
    df.columns = ['SEAS', 'YR', 'ANOM']
    df['SEAS'] = df['SEAS'].astype(str)
    df['YR'] = df['YR'].astype(int)
    df['ANOM'] = df['ANOM'].astype(float)

    # 2. Khởi tạo Workbook & Sheet DATA
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DATA"

    # Hiển thị đường lưới (Gridlines)
    ws.views.sheetView[0].showGridLines = True

    # 3. Định nghĩa Style/Design
    font_family = "Segoe UI"
    
    # Style Header
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Blue
    
    # Style Data
    data_font = Font(name=font_family, size=10, color="000000")
    
    # Alignment
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Border mảnh xám nhạt
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 4. Ghi Header
    headers = ['SEAS', 'YR', 'ANOM']
    ws.append(headers)

    ws.row_dimensions[1].height = 25
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # 5. Ghi dữ liệu & định dạng từng dòng
    for row_idx, row in df.iterrows():
        ws.append([row['SEAS'], row['YR'], row['ANOM']])
        current_row = row_idx + 2  # Dòng dữ liệu bắt đầu từ 2
        
        ws.row_dimensions[current_row].height = 20

        # Cột SEAS (Căn giữa)
        cell_seas = ws.cell(row=current_row, column=1)
        cell_seas.font = data_font
        cell_seas.alignment = align_center
        cell_seas.border = thin_border

        # Cột YR (Căn giữa, định dạng số nguyên)
        cell_yr = ws.cell(row=current_row, column=2)
        cell_yr.font = data_font
        cell_yr.alignment = align_center
        cell_yr.number_format = '0'
        cell_yr.border = thin_border

        # Cột ANOM (Căn phải, định dạng 2 chữ số thập phân)
        cell_anom = ws.cell(row=current_row, column=3)
        cell_anom.font = data_font
        cell_anom.alignment = align_right
        cell_anom.number_format = '0.00'
        cell_anom.border = thin_border

    # 6. Tự động điều chỉnh độ rộng cột (Auto-fit column width)
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

    # 7. Lưu file Excel
    wb.save(output_filename)
    print(f"Đã tạo file Excel thành công: {output_filename}")

# Cách sử dụng:
# convert_txt_to_sea_level_press_excel("D:/Download/olr.txt")
process_relative_oni_to_excel("D:/Download/RONI.ascii.txt")