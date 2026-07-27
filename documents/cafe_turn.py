from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Font, Alignment

def split_cafe_data_to_turns(input_excel_path, output_excel_path):
    """
    Hàm đọc file Excel chứa dữ liệu gộp, tự động nhận diện từng khối cây,
    tách dữ liệu các lần đếm (Lần 1 và Lần 2) sang 2 sheet riêng biệt
    nhưng giữ nguyên cấu trúc đối sánh song song giữa 2 niên vụ.
    """
    wb_in = load_workbook(input_excel_path)
    ws_in = wb_in.active
    
    # --- BƯỚC 1: ĐỌC VÀ PHÂN CỤM DỮ LIỆU THEO TỪNG CÂY ---
    tree_blocks = []
    current_block = []
    current_tree_idx = None
    
    # Quét từ hàng số 3 đến hết file
    for row_idx in range(3, ws_in.max_row + 1):
        tree_val = ws_in.cell(row=row_idx, column=8).value
        # Nếu gặp ô có số hiệu cây -> Đánh dấu bắt đầu một cây mới
        if tree_val is not None and str(tree_val).strip() != "":
            if current_block:
                tree_blocks.append((current_tree_idx, current_block))
            current_block = []
            current_tree_idx = tree_val
            
        if current_tree_idx is not None:
            row_data = {
                'canh_2025': ws_in.cell(row=row_idx, column=1).value,
                'chum_2025': ws_in.cell(row=row_idx, column=2).value,
                'trai_2025': ws_in.cell(row=row_idx, column=3).value,
                'canh_2026': ws_in.cell(row=row_idx, column=5).value,
                'chum_2026': ws_in.cell(row=row_idx, column=6).value,
                'trai_2026': ws_in.cell(row=row_idx, column=7).value,
            }
            current_block.append(row_data)
            
    if current_block:
        tree_blocks.append((current_tree_idx, current_block))

    # --- BƯỚC 2: KHỞI TẠO FILE EXCEL MỚI VÀ 2 SHEET TIÊU ĐỀ ---
    wb_out = Workbook()
    ws1 = wb_out.active
    ws1.title = "Lần 1"
    ws2 = wb_out.create_sheet(title="Lần 2")
    
    def init_sheet_headers(ws):
        # Định dạng tiêu đề lớn cho Niên độ nhóm cột
        ws.merge_cells('A1:C1')
        ws['A1'] = "2025/26"
        ws.merge_cells('E1:G1')
        ws['E1'] = "2026/27"
        
        # Tiêu đề cột nhỏ
        headers = ["canh", "chum", "trai", "", "canh", "chum", "trai", "cây"]
        ws.append(headers)
        
        # Apply Font và Căn giữa cho Tiêu đề
        header_font = Font(name='Arial', size=11, bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        
        ws['A1'].font = header_font; ws['A1'].alignment = center_align
        ws['E1'].font = header_font; ws['E1'].alignment = center_align
        
        for col_idx in range(1, 9):
            cell = ws.cell(row=2, column=col_idx)
            cell.font = header_font
            cell.alignment = center_align

    init_sheet_headers(ws1)
    init_sheet_headers(ws2)

    # --- BƯỚC 3: HÀM PHỤ ĐỂ GHI DỮ LIỆU ĐỒNG BỘ CỦA MỘT CÂY XUỐNG SHEET ---
    def write_tree_to_sheet(ws, tree_idx, rows_2025, rows_2026, start_row):
        max_rows = max(len(rows_2025), len(rows_2026))
        if max_rows == 0:
            return start_row
            
        thin_border = Side(border_style="thin", color="D3D3D3")
        thick_bottom = Side(border_style="medium", color="000000")
        
        current_row = start_row
        for idx in range(max_rows):
            # Điền dữ liệu niên vụ 2025/26
            if idx < len(rows_2025):
                ws.cell(row=current_row, column=1, value=rows_2025[idx]['canh'])
                ws.cell(row=current_row, column=2, value=rows_2025[idx]['chum'])
                ws.cell(row=current_row, column=3, value=rows_2025[idx]['trai'])
            
            ws.cell(row=current_row, column=4, value="")
            
            # Điền dữ liệu niên vụ 2026/27
            if idx < len(rows_2026):
                ws.cell(row=current_row, column=5, value=rows_2026[idx]['canh'])
                ws.cell(row=current_row, column=6, value=rows_2026[idx]['chum'])
                ws.cell(row=current_row, column=7, value=rows_2026[idx]['trai'])
                
            # Điền số hiệu cây ở cột H (chỉ hiển thị ở dòng đầu tiên của cụm cây trong sheet)
            ws.cell(row=current_row, column=8, value=tree_idx if idx == 0 else "")
            
            # Định dạng hàng dữ liệu lưới mỏng
            for col_num in [1, 2, 3, 5, 6, 7, 8]:
                c = ws.cell(row=current_row, column=col_num)
                c.alignment = Alignment(horizontal='center')
                c.border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
                
            # Tạo đường kẻ đậm dưới cùng để phân cách các cây
            if idx == max_rows - 1:
                for col_num in range(1, 9):
                    c = ws.cell(row=current_row, column=col_num)
                    c.border = Border(left=c.border.left, right=c.border.right, top=c.border.top, bottom=thick_bottom)
                    
            current_row += 1
        return current_row

    # --- BƯỚC 4: TÁCH LẦN ĐẾM THEO LOGIC VÀ PHÂN PHỐI VÀO CÁC SHEET ---
    row_sheet1 = 3
    row_sheet2 = 3
    
    for tree_idx, block in tree_blocks:
        # 4.1 Loại bỏ các dòng đệm trống độc lập cho năm 2025
        rows_2025 = []
        for r in block:
            if (r['canh_2025'] is not None and str(r['canh_2025']).strip() != "") or \
               (r['chum_2025'] is not None and str(r['chum_2025']).strip() != "") or \
               (r['trai_2025'] is not None and str(r['trai_2025']).strip() != ""):
                rows_2025.append({'canh': r['canh_2025'], 'chum': r['chum_2025'], 'trai': r['trai_2025']})
                
        # 4.2 Loại bỏ các dòng đệm trống độc lập cho năm 2026
        rows_2026 = []
        for r in block:
            if (r['canh_2026'] is not None and str(r['canh_2026']).strip() != "") or \
               (r['chum_2026'] is not None and str(r['chum_2026']).strip() != "") or \
               (r['trai_2026'] is not None and str(r['trai_2026']).strip() != ""):
                rows_2026.append({'canh': r['canh_2026'], 'chum': r['chum_2026'], 'trai': r['trai_2026']})
                
        # 4.3 Phân tách Lần 1 & Lần 2 của năm 2025 dựa vào số lượng ô có dữ liệu ở cột chum
        rows_2025_l1 = []
        rows_2025_l2 = []
        chum_count_2025 = 0
        for r in rows_2025:
            if r['chum'] is not None and str(r['chum']).strip() != "":
                chum_count_2025 += 1
            if chum_count_2025 <= 2:  # Cặp Chùm-Trái thứ nhất (chứa tối đa 2 ô điền số đếm đầu ở cột chum)
                rows_2025_l1.append(r)
            else:
                rows_2025_l2.append(r)
                
        # 4.4 Phân tách Lần 1 & Lần 2 của năm 2026
        rows_2026_l1 = []
        rows_2026_l2 = []
        chum_count_2026 = 0
        for r in rows_2026:
            if r['chum'] is not None and str(r['chum']).strip() != "":
                chum_count_2026 += 1
            if chum_count_2026 <= 2:
                rows_2026_l1.append(r)
            else:
                rows_2026_l2.append(r)
                
        # 4.5 Ghi dữ liệu vào Sheet 'Lần 1'
        row_sheet1 = write_tree_to_sheet(ws1, tree_idx, rows_2025_l1, rows_2026_l1, row_sheet1)
        
        # 4.6 Ghi dữ liệu vào Sheet 'Lần 2' nếu cây này phát sinh dữ liệu đếm lần 2 ở bất kỳ năm nào
        if len(rows_2025_l2) > 0 or len(rows_2026_l2) > 0:
            row_sheet2 = write_tree_to_sheet(ws2, tree_idx, rows_2025_l2, rows_2026_l2, row_sheet2)
            
    # --- BƯỚC 5: TỰ ĐỘNG CĂN CHỈNH ĐỘ RỘNG CỘT CHO CẢ 2 SHEET ---
    from openpyxl.utils import get_column_letter
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 10)
            
    wb_out.save(output_excel_path)
    print(f"Đã tách thành công thành 2 sheet! File mới lưu tại: {output_excel_path}")


input_file = "D:/Download/output_cafe_Kon_Tum_Mang_Den.xlsx"
output_file = "D:/Download/output_cafe_Kon_Tum_Mang_Den_v2.xlsx"

split_cafe_data_to_turns(input_file, output_file)