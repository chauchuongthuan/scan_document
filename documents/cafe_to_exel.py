import re
import docx
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Font, Alignment

def parse_and_export_cafe_data(docx_path, output_excel_path):
    """
    Hàm đọc dữ liệu từ file Word (.docx), chuẩn hóa cấu trúc theo từng cây,
    chia niên vụ trước/sau dòng '2026-27', bóc tách cột cành linh hoạt
    và xuất ra file Excel chuẩn định dạng.
    """
    # --- BƯỚC 1: ĐỌC FILE WORD VÀ TÁCH THÀNH TỪNG CÂY (Dựa vào khoảng cách hàng) ---
    doc = docx.Document(docx_path)
    
    trees_raw_lines = []
    current_tree_lines = []
    
    for p in doc.paragraphs:
        text = p.text.strip()
        
        # Nếu gặp hàng trống hoặc dòng phân tách dạng gạch ngang '--' -> Kết thúc 1 cây
        if not text or re.match(r'^[-—–_]+$', text):
            if current_tree_lines:
                trees_raw_lines.append(current_tree_lines)
                current_tree_lines = []
        else:
            current_tree_lines.append(text)
            
    # Đóng cụm cây cuối cùng nếu còn sót
    if current_tree_lines:
        trees_raw_lines.append(current_tree_lines)

    # --- BƯỚC 2: PHÂN TÍCH LOGIC NIÊN VỤ & CỘT CÀNH CHO TỪNG CÂY ---
    processed_trees = []
    
    for tree_lines in trees_raw_lines:
        # Khởi tạo cấu trúc chứa dữ liệu của cây hiện tại cho 2 năm
        tree_data = {
            '2025/26': {'canh': None, 'lines': []},
            '2026/27': {'canh': None, 'lines': []}
        }
        
        # Mặc định ban đầu luôn là năm 2025/26 cho đến khi gặp dòng "2026-27"
        current_year = '2025/26'
        
        for line in tree_lines:
            # Nếu gặp dòng đánh dấu năm -> Chuyển trạng thái sang năm 2026/27
            if '2026-27' in line:
                current_year = '2026/27'
                continue
            
            # Kiểm tra dòng ghi nhận "cành" (chấp nhận gõ sai chính tả 'cành'/'càng')
            canh_match = re.search(r'(\d+)\s*càn[hg]', line, re.IGNORECASE)
            if canh_match:
                tree_data[current_year]['canh'] = int(canh_match.group(1))
                continue
            
            # Kiểm tra dòng dữ liệu số (bắt đầu bằng số và ký tự 'c')
            if re.match(r'^\d+c', line):
                tree_data[current_year]['lines'].append(line)
                
        processed_trees.append(tree_data)

    # --- BƯỚC 3: HÀM PHỤ ĐỂ PHÁT TRIỂN CẶP DỮ LIỆU (CHÙM - TRÁI) THÀNH CÁC DÒNG EXCEL ---
    def expand_year_to_rows(year_structure):
        lines = year_structure['lines']
        canh_val = year_structure['canh']
        expanded_rows = []
        
        # Duyệt qua các dòng dữ liệu theo từng cặp (Dòng 1: Chùm, Dòng 2: Trái)
        for i in range(0, len(lines), 2):
            chum_line = lines[i]
            trai_line = lines[i+1] if (i + 1) < len(lines) else ""
            
            # Phân tách số lượng chùm và danh sách số đi kèm
            chum_match = re.match(r'^(\d+)c\s*(.*)$', chum_line)
            chum_prefix = int(chum_match.group(1)) if chum_match else ""
            chum_vals = [int(x) for x in re.findall(r'\d+', chum_match.group(2))] if chum_match else []
            
            # Phân tách số lượng trái và danh sách số đi kèm
            trai_prefix = ""
            trai_vals = []
            if trai_line:
                trai_match = re.match(r'^(\d+)c\s*(.*)$', trai_line)
                trai_prefix = int(trai_match.group(1)) if trai_match else ""
                trai_vals = [int(x) for x in re.findall(r'\d+', trai_match.group(2))] if trai_match else []
            
            # Gom toàn bộ các hạt/trái của cặp này vào chung 1 cột dọc (trái)
            all_trai_values = chum_vals + trai_vals
            num_rows_for_pair = max(2, len(all_trai_values))
            
            for r in range(num_rows_for_pair):
                chum_cell = ""
                if r == 0:
                    chum_cell = chum_prefix
                elif r == 1:
                    chum_cell = trai_prefix
                    
                trai_cell = all_trai_values[r] if r < len(all_trai_values) else ""
                
                expanded_rows.append({'chum': chum_cell, 'trai': trai_cell})
                
        # Tạo cấu trúc dòng cuối cùng bao gồm cả cột 'cành' (chỉ xuất hiện ở dòng đầu tiên của khối năm)
        total_rows_needed = max(1 if canh_val is not None else 0, len(expanded_rows))
        final_rows = []
        for r in range(total_rows_needed):
            final_rows.append({
                'canh': canh_val if r == 0 else "",
                'chum': expanded_rows[r]['chum'] if r < len(expanded_rows) else "",
                'trai': expanded_rows[r]['trai'] if r < len(expanded_rows) else ""
            })
        return final_rows

    # --- BƯỚC 4: KHỞI TẠO FILE EXCEL VÀ GHI DỮ LIỆU TIÊU ĐỀ ---
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Cafe"
    
    # Định dạng tiêu đề lớn cho Niên độ nhóm cột
    ws.merge_cells('A1:C1')
    ws['A1'] = "2025/26"
    ws.merge_cells('E1:G1')
    ws['E1'] = "2026/27"
    
    # Tiêu đề cột nhỏ (Bổ sung thêm cột 'cây' ở vị trí thứ 8)
    headers = ["canh", "chum", "trai", "", "canh", "chum", "trai", "cây"]
    ws.append(headers)
    
    # Apply Font và Căn giữa cho Tiêu đề
    header_font = Font(name='Arial', size=11, bold=True)
    center_align = Alignment(horizontal='center', vertical='center')
    
    ws['A1'].font = header_font; ws['A1'].alignment = center_align
    ws['E1'].font = header_font; ws['E1'].alignment = center_align
    
    # Mở rộng range từ (1, 8) lên (1, 9) để định dạng thêm cột thứ 8 (cột H)
    for col_idx in range(1, 9):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.alignment = center_align

    # --- BƯỚC 5: ĐIỀN DỮ LIỆU ĐỒNG BỘ SONG SONG THEO TỪNG CÂY ---
    current_row = 3
    thin_border = Side(border_style="thin", color="D3D3D3")
    thick_bottom = Side(border_style="medium", color="000000") # Viền đậm phân tách cây
    
    # Sử dụng enumerate để lấy số thứ tự cây (bắt đầu từ 1)
    for tree_idx, tree in enumerate(processed_trees, start=1):
        rows_2025 = expand_year_to_rows(tree['2025/26'])
        rows_2026 = expand_year_to_rows(tree['2026/27'])
        
        # Số lượng dòng tối đa mà cây này chiếm dụng trong Excel để giữ cân bằng hàng ngang
        max_tree_rows = max(len(rows_2025), len(rows_2026))
        
        for idx in range(max_tree_rows):
            # Điền dữ liệu năm 2025/26 (Cột A, B, C)
            if idx < len(rows_2025):
                ws.cell(row=current_row, column=1, value=rows_2025[idx]['canh'])
                ws.cell(row=current_row, column=2, value=rows_2025[idx]['chum'])
                ws.cell(row=current_row, column=3, value=rows_2025[idx]['trai'])
            
            # Cột D để trống làm khoảng cách ngăn cách giữa 2 năm
            ws.cell(row=current_row, column=4, value="")
            
            # Điền dữ liệu năm 2026/27 (Cột E, F, G)
            if idx < len(rows_2026):
                ws.cell(row=current_row, column=5, value=rows_2026[idx]['canh'])
                ws.cell(row=current_row, column=6, value=rows_2026[idx]['chum'])
                ws.cell(row=current_row, column=7, value=rows_2026[idx]['trai'])
            
            # Điền số thứ tự cây vào cột thứ 8 (Cột H) - Chỉ ghi ở dòng đầu tiên của cụm cây
            ws.cell(row=current_row, column=8, value=tree_idx if idx == 0 else "")
            
            # Định dạng căn lề và đường lưới cho hàng hiện tại (Bổ sung thêm cột số 8)
            for col_num in [1, 2, 3, 5, 6, 7, 8]:
                c = ws.cell(row=current_row, column=col_num)
                c.alignment = Alignment(horizontal='center')
                c.border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
            
            # Nếu là dòng cuối cùng của CÂY hiện tại -> Đánh đường viền gạch dưới đậm để phân tách với cây sau
            if idx == max_tree_rows - 1:
                for col_num in range(1, 9):
                    c = ws.cell(row=current_row, column=col_num)
                    c.border = Border(left=c.border.left, right=c.border.right, top=c.border.top, bottom=thick_bottom)
                    
            current_row += 1

    # Tự động căn chỉnh độ rộng các cột cho đẹp mắt
    from openpyxl.utils import get_column_letter
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 10)
        
    wb.save(output_excel_path)
    print(f"Xử lý thành công! File đã lưu tại: {output_excel_path}")

# --- CÁCH SỬ DỤNG HÀM ---
parse_and_export_cafe_data("D:/Download/cafe_Kon_Tum_Mang_Den.docx", "D:/Download/output_cafe_Kon_Tum_Mang_Den.xlsx")