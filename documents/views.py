from django.http import JsonResponse, HttpResponse
from django.shortcuts import render
from django.shortcuts import render, get_object_or_404, redirect
from django.core.paginator import Paginator
from django.db.models import Q
from django.contrib import messages
from django.urls import reverse
from django.utils.text import slugify

from coffee.upload import pathByDate, uploadFile, upload_editor_image
from .models import Document
from django.utils import timezone
import os
import urllib.parse
from django.core.files.storage import default_storage
from django.contrib.auth.decorators import login_required
import pandas as pd
import json
from coffee.helpers import extract_fields_from_pdf
import openpyxl
# from openpyxl.styles import Alignment
import datetime
# Document Views
class DocumentListView:
    @login_required(login_url="login_cms")
    def list_documents(request):
        per_page_options = [10, 25, 50, 100]

        per_page = request.GET.get('per_page', 10)

        query_name_search = request.GET.get('name_search')
        documents = Document.objects.all().order_by('-created_at')
        
        if query_name_search:
            documents = documents.filter(
                Q(comco_ref__icontains=query_name_search) |
                Q(buyer_ref__icontains=query_name_search)
            )
            
        paginator = Paginator(documents, per_page)
        page_number = request.GET.get('page')
        page_obj = paginator.get_page(page_number)

        return render(request, 'documents/list.html', {'documents': page_obj, 'per_page_options': per_page_options, 'per_page': int(per_page),})
    
    @login_required(login_url="login_cms")
    def document_detail(request, code):
        document = get_object_or_404(Document, code=code)
        return render(request, 'documents/document_detail.html', {'document': document})
    
    @login_required(login_url="login_cms")
    def create_document(request):
        if request.method == 'POST':
            try:
                # Lấy dữ liệu từ request.POST
                comco_ref = request.POST.get('comco_ref')
                buyer_ref = request.POST.get('buyer_ref')
                contract_date = request.POST.get('contract_date')
                price = request.POST.get('price')
                dds_and_geo = request.POST.get('dds_and_geo')
                docs = request.POST.get('docs')
                shippers = request.POST.get('shippers')
                buyers = request.POST.get('buyers')
                shipment_period = request.POST.get('shipment_period')
                quality = request.POST.get('quality')
                quantity = request.POST.get('quantity')
                conts_no = request.POST.get('conts_no')
                packing = request.POST.get('packing')
                sample = request.POST.get('sample')
                result = request.POST.get('result')
                stuffing_date = request.POST.get('stuffing_date') or None
                carriers = request.POST.get('carriers')
                booking_no = request.POST.get('booking_no')
                etd = request.POST.get('etd') or None
                eta = request.POST.get('eta')
                destination = request.POST.get('destination')
                inspectors = request.POST.get('inspectors')
                warehouse = request.POST.get('warehouse')
                status = request.POST.get('status')
                weight_claim = request.POST.get('weight_claim')
                remark = request.POST.get('remark')
                
                document = Document(
                    comco_ref=comco_ref,
                    buyer_ref=buyer_ref,
                    contract_date=contract_date,
                    price=price,
                    dds_and_geo=dds_and_geo,
                    docs=docs,
                    shippers=shippers,
                    buyers=buyers,
                    shipment_period=shipment_period,
                    quality=quality,
                    quantity=quantity,
                    conts_no=conts_no,
                    packing=packing,
                    sample=sample,
                    result=result,
                    stuffing_date=stuffing_date,
                    carriers=carriers,
                    booking_no=booking_no,
                    etd=etd,
                    eta=eta,
                    destination=destination,
                    inspectors=inspectors,
                    warehouse=warehouse,
                    status=status,
                    weight_claim=weight_claim,
                    remark=remark,
                )
                document.save()
                
                return JsonResponse({
                    'status': 'success',
                    'message': 'Thêm hợp đồng thành công!',
                    'redirect': reverse('document-list')
                })
            except Exception as e:
                print("Error:", str(e))  # Debug log
                return JsonResponse({
                    'status': 'error',
                    'message': f'Lỗi khi thêm hợp đồng: {str(e)}'
                }, status=400)  
        
        return render(request, 'documents/create.html')
    
    @login_required(login_url="login_cms")
    def import_excel(request):
        if request.method == 'POST' and request.FILES.get('excel_file'):
            try:
                excel_file = request.FILES['excel_file']

                # Save the uploaded file temporarily
                temp_file_path = default_storage.save('temp/' + excel_file.name, excel_file)
                file_path = os.path.join(default_storage.location, temp_file_path)

                # Read Excel file
                df = pd.read_excel(
                    file_path,
                    header=2,
                    parse_dates=['CONTRACT DATE', 'STUFFING DATE', 'ETD'], # Specify columns to parse
                    date_parser=lambda x: pd.to_datetime(x, format='%d/%m/%Y', errors='coerce')
                )

                # Remove leading/trailing whitespace and convert to uppercase
                df.columns = df.columns.str.strip().str.upper()
                
                # Expected columns based on Document model
                expected_columns = [
                    'COMCO REF', 'BUYER REF', 'CONTRACT DATE', 'PRICE', 'DDS & GEO',
                    'DOCS', 'SHIPPERS', 'BUYERS', 'SHIPMENT PERIOD', 'QUALITY',
                    'QUANTITY', 'CONTS NO', 'PACKING', 'SAMPLE', 'RESULTS',
                    'STUFFING DATE', 'CARRIERS', 'BOOKING NO', 'ETD', 'ETA',
                    'DESTINATION', 'INSPECTORS', 'WAREHOUSE', 'STATUS',
                    'WEIGHT CLAIM', 'REMARK'
                ]

                # Check if all required columns are present
                missing_columns = [col for col in expected_columns if col not in df.columns]
                if missing_columns:
                    return JsonResponse({
                        'status': 'error',
                        'message': f'Các cột thiếu trong file Excel: {", ".join(missing_columns)}'
                    }, status=400)

                imported_count = 0
                error_count = 0
                errors = []

                for index, row in df.iterrows():
                    try:
                        def safe_str_excel(value):
                            return "" if pd.isna(value) else str(value)
                        
                        document = Document(
                            comco_ref=safe_str_excel(row.get('COMCO REF', '')),
                            buyer_ref=safe_str_excel(row.get('BUYER REF', '')),
                            contract_date=row.get('CONTRACT DATE'),
                            price=safe_str_excel(row.get('PRICE', '')),
                            dds_and_geo=safe_str_excel(row.get('DDS & GEO', '')),
                            docs=safe_str_excel(row.get('DOCS', '')),
                            shippers=safe_str_excel(row.get('SHIPPERS', '')),
                            buyers=safe_str_excel(row.get('BUYERS', '')),
                            shipment_period=safe_str_excel(row.get('SHIPMENT PERIOD', '')),
                            quality=safe_str_excel(row.get('QUALITY', '')),
                            quantity=safe_str_excel(row.get('QUANTITY', '')),
                            conts_no=safe_str_excel(row.get('CONTS NO', '')),
                            packing=safe_str_excel(row.get('PACKING', '')),
                            sample=safe_str_excel(row.get('SAMPLE', '')),
                            result=safe_str_excel(row.get('RESULT', '')),
                            stuffing_date=row.get('STUFFING DATE'),
                            carriers=safe_str_excel(row.get('CARRIERS', '')),
                            booking_no=safe_str_excel(row.get('BOOKING NO', '')),
                            etd=row.get('ETD'),
                            eta=safe_str_excel(row.get('ETA', '')),
                            destination=safe_str_excel(row.get('DESTINATION', '')),
                            inspectors=safe_str_excel(row.get('INSPECTORS', '')),
                            warehouse=safe_str_excel(row.get('WAREHOUSE', '')),
                            status=safe_str_excel(row.get('STATUS', '')),
                            weight_claim=safe_str_excel(row.get('WEIGHT CLAIM', '')),
                            remark=safe_str_excel(row.get('REMARK', ''))
                        )
                        document.save()
                        imported_count += 1

                    except Exception as e:
                        error_count += 1
                        errors.append(f"Dòng {index + 1}") 

                
                os.remove(file_path)    
                
                message = f'Import thành công {imported_count} bản ghi. {error_count} bản ghi lỗi: {errors}'

                if error_count > 0 and error_count < len(df):
                    return JsonResponse({
                        'status': 'warning',
                        'message': message,
                        'imported_count': imported_count,
                        'error_count': error_count,
                        'errors': errors
                    })
                elif error_count > 0 and error_count == len(df):
                    return JsonResponse({
                        'status': 'error',
                        'message': 'Import thất bại, tất cả các bản ghi đều lỗi',
                        'imported_count': imported_count,
                        'error_count': error_count,
                        'errors': errors
                    })
                elif error_count == 0:
                    return JsonResponse({
                        'status': 'success',
                        'message': message,
                        'imported_count': imported_count,
                        'error_count': error_count,
                        'errors': errors
                })

            except Exception as e:
                return JsonResponse({
                    'status': 'error',
                    'message': f'Lỗi khi import file Excel: {str(e)}'
                }, status=400)

        return JsonResponse({
            'status': 'error',
            'message': 'Phương thức không được hỗ trợ hoặc thiếu file Excel'
        }, status=405)

    @login_required(login_url="login_cms")
    def export_excel(request):
        try:
            # Lấy tất cả documents từ database
            documents = Document.objects.all().order_by('comco_ref')
            
            # Tạo DataFrame từ dữ liệu documents
            data = []
            for document in documents:
                # Format date fields to dd/mm/yyyy
                def format_date(date_field):
                    if date_field:
                        if isinstance(date_field, (datetime.date, datetime.datetime)):
                            return date_field.strftime('%d/%m/%Y')
                        # Nếu là string, cố gắng parse và format lại
                        try:
                            parsed_date = datetime.datetime.strptime(str(date_field), '%Y-%m-%d')
                            return parsed_date.strftime('%d/%m/%Y')
                        except:
                            return str(date_field)
                    return ""
                
                data.append({
                    'COMCO REF': document.comco_ref,
                    'BUYER REF': document.buyer_ref,
                    'CONTRACT DATE': format_date(document.contract_date),
                    'PRICE': document.price,
                    'DDS & GEO': document.dds_and_geo,
                    'DOCS': document.docs,
                    'SHIPPERS': document.shippers,
                    'BUYERS': document.buyers,
                    'SHIPMENT PERIOD': document.shipment_period,
                    'QUALITY': document.quality,
                    'QUANTITY': document.quantity,
                    'CONTS NO': document.conts_no,
                    'PACKING': document.packing,
                    'SAMPLE': document.sample,
                    'RESULTS': document.result,
                    'STUFFING DATE': format_date(document.stuffing_date),
                    'CARRIERS': document.carriers,
                    'BOOKING NO': document.booking_no,
                    'ETD': format_date(document.etd),
                    'ETA': document.eta,
                    'DESTINATION': document.destination,
                    'INSPECTORS': document.inspectors,
                    'WAREHOUSE': document.warehouse,
                    'STATUS': document.status,
                    'WEIGHT CLAIM': document.weight_claim,
                    'REMARK': document.remark
                })
            
            df = pd.DataFrame(data)
            # Tạo HTTP response
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            filename = f"schedule_export.xlsx"
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            
            # Tạo Excel file với cấu trúc giống file mẫu
            with pd.ExcelWriter(response, engine='openpyxl') as writer:
                # Tạo workbook
                workbook = writer.book
                
                worksheet = workbook.create_sheet("SCHEDULE", 0)
                
                # Xóa sheet mặc định nếu có
                if 'Sheet' in workbook.sheetnames:
                    del workbook['Sheet']
                
                # Merge dòng 1 và dòng 2: Tiêu đề "SCHEDULE" (merge từ A1 đến Z2)
                worksheet.merge_cells('A1:Z2')
                worksheet['A1'] = 'SCHEDULE'
                
                # Dòng 3: Headers
                headers = [
                    'COMCO REF', 'BUYER REF', 'CONTRACT DATE', 'PRICE', 'DDS & GEO', 'DOCS',
                    'SHIPPERS', 'BUYERS', 'SHIPMENT PERIOD', 'QUALITY', 'QUANTITY', 'CONTS NO',
                    'PACKING', 'SAMPLE', 'RESULTS', 'STUFFING DATE', 'CARRIERS', 'BOOKING NO',
                    'ETD', 'ETA', 'DESTINATION', 'INSPECTORS', 'WAREHOUSE', 'STATUS',
                    'WEIGHT CLAIM', 'REMARK'
                ]
                
                for col, header in enumerate(headers, 1):
                    cell = worksheet.cell(row=3, column=col, value=header)
                    # Có thể thêm style cho header nếu cần
                    cell.font = openpyxl.styles.Font(bold=True)
                
                # Ghi dữ liệu từ DataFrame (bắt đầu từ dòng 4)
                if not df.empty:
                    for row_idx, row_data in enumerate(df.itertuples(), 4):
                        for col_idx, value in enumerate(row_data[1:], 1):  # Bỏ qua index
                            worksheet.cell(row=row_idx, column=col_idx, value=value)
            
            return response
            
        except Exception as e:
            return JsonResponse({
                'status': 'error',
                'message': f'Export Excel thất bại: {str(e)}'
            }, status=400)

    @login_required(login_url="login_cms")
    def scan_pdf(request):
        if request.method == 'POST' and request.FILES.get('pdf_file'):
            try:
                pdf_file = request.FILES['pdf_file']
                form_type = request.POST.get('form_type')
                extracted_data = extract_fields_from_pdf(pdf_file, form_type)
                
                print("extracted_data", extracted_data)
                return JsonResponse({
                    'status': 'success',
                    'message': 'Scan File thành công!',
                    'data': extracted_data
                })
                
            except Exception as e:
                print("e", e)
                return JsonResponse({
                    'status': 'error',
                    'message': f'Lỗi khi scan file: {str(e)}'
                }, status=400)

        return JsonResponse({
            'status': 'error',
            'message': 'Phương thức không được hỗ trợ hoặc thiếu file'
        }, status=405)

    @login_required(login_url="login_cms")
    def update_document(request, pk):
        document = get_object_or_404(Document, pk=pk)
        if request.method == 'POST':
            try:               
                comco_ref = request.POST.get('comco_ref')
                buyer_ref = request.POST.get('buyer_ref')
                contract_date = request.POST.get('contract_date')
                price = request.POST.get('price')
                dds_and_geo = request.POST.get('dds_and_geo')
                docs = request.POST.get('docs')
                shippers = request.POST.get('shippers')
                buyers = request.POST.get('buyers')
                shipment_period = request.POST.get('shipment_period')
                quality = request.POST.get('quality')
                quantity = request.POST.get('quantity')
                conts_no = request.POST.get('conts_no')
                packing = request.POST.get('packing')
                sample = request.POST.get('sample')
                result = request.POST.get('result')
                stuffing_date = request.POST.get('stuffing_date') or None
                carriers = request.POST.get('carriers')
                booking_no = request.POST.get('booking_no')
                etd = request.POST.get('etd') or None
                eta = request.POST.get('eta')
                destination = request.POST.get('destination')
                inspectors = request.POST.get('inspectors')
                warehouse = request.POST.get('warehouse')
                status = request.POST.get('status')
                weight_claim = request.POST.get('weight_claim')
                remark = request.POST.get('remark')

                document.comco_ref = comco_ref
                document.buyer_ref = buyer_ref
                document.contract_date = contract_date
                document.price = price
                document.dds_and_geo = dds_and_geo
                document.docs = docs
                document.shippers = shippers
                document.buyers = buyers
                document.shipment_period = shipment_period
                document.quality = quality
                document.quantity = quantity
                document.conts_no = conts_no
                document.packing = packing
                document.sample = sample
                document.result = result
                document.stuffing_date = stuffing_date
                document.carriers = carriers
                document.booking_no = booking_no
                document.etd = etd
                document.eta = eta
                document.destination = destination
                document.inspectors = inspectors
                document.warehouse = warehouse
                document.status = status
                document.weight_claim = weight_claim
                document.remark = remark

                document.save()
                return JsonResponse({
                    'status': 'success',
                    'message': 'Cập nhật hợp đồng thành công!',
                    'redirect': reverse('document-list')
                })
            except Exception as e:
                return JsonResponse({
                    'status': 'error',
                    'message': f'Lỗi khi cập nhật hợp đồng: {str(e)}'
                }, status=400)
        
        return render(request, 'documents/update.html', {'document': document})

    @login_required(login_url="login_cms")
    def delete_document(request, pk):
        document = get_object_or_404(Document, pk=pk)
        if request.method == 'POST':
            try:    
                document.delete()
                return JsonResponse({
                    'status': 'success',
                    'message': 'Xóa hợp đồng thành công!',
                }, status=200)
            except Exception as e:
                return JsonResponse({
                    'status': 'error',
                    'message': f'Lỗi khi xóa hợp đồng: {str(e)}',
                }, status=400)